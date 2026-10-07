import asyncio
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from revit_model_mcp import server
from revit_model_mcp.issue_register import (
    COLORS,
    SEVERITIES,
    snapshot_indices,
    validate_capture,
    validate_register,
    write_register,
)


def issue(**changes):
    return {
        "title": "Missing code",
        "category": "Classification",
        "finding": "2 types have no code",
        "severity": "major",
        **changes,
    }


def test_workbook_sheets_rows_images_and_totals(tmp_path):
    png = tmp_path / "image.png"
    Image.new("RGB", (40, 20), "red").save(png)
    issues = [
        issue(severity=value, element_ids=[index + 1]) for index, value in enumerate(SEVERITIES)
    ]
    output = tmp_path / "register.xlsx"
    result = write_register(
        str(output),
        {
            "name": "Demo",
            "reviewer": "Reviewer",
            "documents": [{"title": "EIR", "reference": "4.3", "revision": "P01"}],
        },
        issues,
        {0: png},
        {1: "Snapshot unavailable: capture failed"},
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Cover", "Summary", "Register", "Elements"]
    assert workbook.properties.title == "Demo issue register"
    assert workbook.properties.creator == "Reviewer"
    assert workbook["Cover"]["A15"].value == "EIR"
    register = workbook["Register"]
    assert register.max_row == 5
    assert register.freeze_panes == "C2"
    assert register.auto_filter.ref == "A1:N5"
    assert all(cell.font.bold and cell.font.color.rgb == "00FFFFFF" for cell in register[1])
    for index, color in enumerate(COLORS, 2):
        assert register.cell(index, 4).fill.fgColor.rgb == f"00{color}"
    assert register["D2"].font.color.rgb == "00FFFFFF"
    assert len(register._images) == 1
    assert register._images[0].anchor._from.col == 1
    assert register._images[0].anchor._from.row == 1
    assert register.row_dimensions[2].height == 185
    assert register.row_dimensions[3].height == 60
    assert register["B3"].value == "Snapshot unavailable: capture failed"
    assert workbook["Elements"].max_row == 5
    assert workbook["Summary"]["F3"].value == 4
    assert len(workbook["Summary"]._charts) == 1
    assert workbook["Summary"]._charts[0].grouping == "stacked"
    assert result["snapshotCount"] == 1
    assert result["bySeverity"] == dict.fromkeys(SEVERITIES, 1)


def test_preserves_existing_output(tmp_path):
    output = tmp_path / "register.xlsx"
    output.write_bytes(b"original")
    with pytest.raises(ValueError, match="output_path already exists"):
        write_register(str(output), {}, [issue()])
    assert output.read_bytes() == b"original"


@pytest.mark.parametrize(
    "issues,field",
    [
        ([issue(title=None)], "issues[0].title"),
        ([issue(severity="high")], "issues[0].severity"),
        ([issue()] * 61, "issues"),
        ([issue(element_ids=[True])], "issues[0].element_ids"),
        ([issue(snapshot="elevation")], "issues[0].snapshot"),
        ([issue(finding="bad\x00text")], "issues[0].finding"),
    ],
)
def test_validation_errors(tmp_path, issues, field):
    with pytest.raises(ValueError) as caught:
        validate_register(str(tmp_path / "register.xlsx"), {}, issues)
    assert field in str(caught.value)


def test_snapshot_cap_prioritizes_severity_and_keeps_order(tmp_path):
    issues = [issue(severity="minor", element_ids=[1]) for _ in range(25)]
    issues.extend(
        [issue(severity="critical", element_ids=[2]), issue(severity="major", element_ids=[3])]
    )
    _, _, issues = validate_register(str(tmp_path / "out.xlsx"), {}, issues)
    selected, skipped = snapshot_indices(issues)
    assert selected == [25, 26, *range(23)]
    assert skipped == [23, 24]
    result = write_register(str(tmp_path / "out.xlsx"), {}, issues)
    register = load_workbook(result["path"])["Register"]
    assert register["B25"].value == "Snapshot skipped: limit of 25 per register"
    assert sum("limit of 25" in warning for warning in result["warnings"]) == 2


def test_invalid_image_becomes_warning(tmp_path):
    png = tmp_path / "bad.png"
    png.write_bytes(b"invalid image")
    result = write_register(str(tmp_path / "out.xlsx"), {}, [issue(element_ids=[1])], {0: png})
    assert result["snapshotCount"] == 0
    assert "Snapshot unavailable:" in result["warnings"][0]


@pytest.mark.parametrize(
    "changes",
    [{"element_ids": []}, {"pixel_size": 4001}, {"padding_mm": float("nan")}, {"mode": "bad"}],
)
def test_capture_validation(changes):
    with pytest.raises(ValueError):
        validate_capture(**{"element_ids": [1], **changes})


def test_tool_captures_sequentially_with_addressing_and_continues_on_error(tmp_path):
    calls = []

    async def capture(job, *_):
        calls.append(job)
        if job.payload["elementIds"] == [2]:
            raise server.ToolError("capture failed")
        Image.new("RGB", (40, 20), "red").save(job.output_path)
        return {"success": True, "data": {"localPath": job.output_path}}

    issues = [
        issue(element_ids=[1]),
        issue(element_ids=[2], snapshot="plan"),
        issue(element_ids=[3], snapshot="none"),
        issue(),
    ]
    with (
        patch.object(server.channel, "execute", side_effect=capture),
        patch.dict("os.environ", {"REVIT_MCP_READ_ONLY": "1"}),
    ):
        result = asyncio.run(
            server.mcp.call_tool(
                "revit_issue_register",
                {
                    "output_path": str(tmp_path / "out.xlsx"),
                    "project": {},
                    "issues": issues,
                    "document": "Demo",
                    "process_id": 42,
                },
            )
        )
    assert not result.is_error
    assert len(calls) == 2
    for job, ids, mode in zip(calls, ([1], [2]), ("3d", "plan"), strict=True):
        assert job.command == "capture-elements"
        assert job.payload == {
            "command": "capture-elements",
            "elementIds": ids,
            "pixelSize": 900,
            "paddingMm": 1500,
            "mode": mode,
            "targetDocument": "Demo",
            "targetProcessId": 42,
        }
    register = load_workbook(tmp_path / "out.xlsx")["Register"]
    assert len(register._images) == 1
    assert register["B3"].value == "Snapshot unavailable: capture failed"
    assert not Path(calls[0].output_path).exists()


def test_tool_validates_all_issues_before_contacting_revit(tmp_path):
    with patch.object(server.channel, "execute", new=AsyncMock()) as execute:
        with pytest.raises(server.ToolError, match=r"issues\[1\].title"):
            asyncio.run(
                server.mcp.call_tool(
                    "revit_issue_register",
                    {
                        "output_path": str(tmp_path / "out.xlsx"),
                        "project": {},
                        "issues": [issue(element_ids=[1]), issue(title="")],
                    },
                )
            )
    execute.assert_not_awaited()


def test_tool_cap_sends_only_25_jobs_in_severity_order(tmp_path):
    calls = []

    async def capture(job, *_):
        calls.append(job.payload["elementIds"][0])
        raise server.ToolError("unavailable")

    issues = [issue(severity="info", element_ids=[index + 1]) for index in range(26)]
    issues.append(issue(severity="critical", element_ids=[100]))
    with patch.object(server.channel, "execute", side_effect=capture):
        result = asyncio.run(
            server.mcp.call_tool(
                "revit_issue_register",
                {"output_path": str(tmp_path / "out.xlsx"), "project": {}, "issues": issues},
            )
        )
    assert not result.is_error
    assert calls == [100, *range(1, 25)]


def test_capture_tool_returns_png_content(tmp_path):
    png = tmp_path / "capture.png"
    Image.new("RGB", (2, 2), "red").save(png)
    with patch.object(
        server.channel,
        "execute",
        new=AsyncMock(return_value={"success": True, "data": {"localPath": str(png)}}),
    ) as execute:
        result = asyncio.run(
            server.mcp.call_tool(
                "revit_capture_elements",
                {"element_ids": [1], "mode": "plan", "padding_mm": 0, "output_path": str(png)},
            )
        )
    assert not result.is_error
    assert any(
        content.type == "image" and content.mime_type == "image/png" for content in result.content
    )
    assert execute.await_args.args[0].payload["mode"] == "plan"


def test_capture_budget_preserves_register_and_skips_remaining(tmp_path):
    calls = []

    async def slow_capture(job, *_):
        calls.append(job)
        await asyncio.sleep(1)

    with (
        patch.object(server, "ISSUE_CAPTURE_BUDGET_SECONDS", 0.01),
        patch.object(server.channel, "execute", side_effect=slow_capture),
    ):
        result = asyncio.run(
            server.revit_issue_register(
                str(tmp_path / "out.xlsx"),
                {},
                [issue(element_ids=[1]), issue(element_ids=[2])],
            )
        )
    assert len(calls) == 1
    assert result["snapshotCount"] == 0
    assert len(result["warnings"]) == 2
    register = load_workbook(result["path"])["Register"]
    assert register["B2"].value == "Snapshot unavailable: capture time budget exhausted"
    assert register["B3"].value == "Snapshot skipped: capture time budget exhausted"


def test_early_budget_timer_starts_one_capture(tmp_path):
    calls = []

    async def run():
        async def stuck_capture(job, *_):
            calls.append(job)
            await asyncio.Event().wait()

        with (
            patch.object(server, "ISSUE_CAPTURE_BUDGET_SECONDS", 0.2),
            patch.object(server.channel, "execute", side_effect=stuck_capture),
        ):
            return await server.revit_issue_register(
                str(tmp_path / "out.xlsx"),
                {},
                [issue(element_ids=[1]), issue(element_ids=[2]), issue(element_ids=[3])],
            )

    loop = asyncio.new_event_loop()
    # asyncio fires timers up to this much early, as on Windows (15.6 ms).
    loop._clock_resolution = 1.0
    try:
        result = loop.run_until_complete(run())
    finally:
        loop.close()
    assert len(calls) == 1
    register = load_workbook(result["path"])["Register"]
    assert register["B2"].value == "Snapshot unavailable: capture time budget exhausted"
    assert register["B3"].value == "Snapshot skipped: capture time budget exhausted"
    assert register["B4"].value == "Snapshot skipped: capture time budget exhausted"


def test_output_validation_before_revit(tmp_path):
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"original")
    with patch.object(server.channel, "execute", new=AsyncMock()) as execute:
        with pytest.raises(server.ToolError, match="output_path already exists"):
            asyncio.run(server.revit_issue_register(str(output), {}, [issue(element_ids=[1])]))
    execute.assert_not_awaited()
    assert output.read_bytes() == b"original"


def test_workbook_text_is_literal(tmp_path):
    output = tmp_path / "literal.xlsx"
    write_register(str(output), {"name": "=project"}, [issue(requirement="=1+1")])
    workbook = load_workbook(output)
    assert workbook["Cover"]["C6"].value == "=project"
    assert workbook["Cover"]["C6"].data_type == "s"
    assert workbook["Register"]["G2"].value == "=1+1"
    assert workbook["Register"]["G2"].data_type == "s"


def test_register_column_layout_and_full_element_inventory(tmp_path):
    output = tmp_path / "register.xlsx"
    write_register(str(output), {}, [issue(element_ids=list(range(1, 36)), status="Closed")])
    workbook = load_workbook(output)
    register = workbook["Register"]
    assert [cell.value for cell in register[1]] == [
        "ID",
        "Snapshot",
        "Title",
        "Severity",
        "Category",
        "Requirement source",
        "Requirement",
        "Finding",
        "Recommendation",
        "Status",
        "Responsible",
        "Due",
        "Elements",
        "Element IDs",
    ]
    assert [register.column_dimensions[cell.column_letter].width for cell in register[1]] == [
        9,
        52,
        32,
        11,
        20,
        16,
        40,
        46,
        36,
        10,
        16,
        11,
        9,
        24,
    ]
    assert register["C2"].value == "Missing code"
    assert register["H2"].value == "2 types have no code"
    assert register["J2"].value == "Closed"
    assert register["J2"].fill.fgColor.rgb == "00F2F2F2"
    assert register["M2"].value == 35
    assert register["N2"].value == ", ".join(map(str, range(1, 31))) + " +5 more"
    assert workbook["Elements"].max_row == 36
    assert workbook["Elements"]["B36"].value == 35


def test_tall_snapshot_is_trimmed_boxed_and_anchored_inside_cell(tmp_path):
    png = tmp_path / "tall.png"
    source = Image.new("RGB", (600, 1200), (250, 250, 250))
    ImageDraw.Draw(source).rectangle((270, 200, 329, 999), fill="red")
    source.save(png)
    output = tmp_path / "register.xlsx"
    write_register(str(output), {}, [issue(element_ids=[1])], {0: png})
    with ZipFile(output) as archive:
        with Image.open(BytesIO(archive.read("xl/media/image1.png"))) as boxed:
            assert boxed.size == (360, 240)
            assert boxed.getpixel((0, 0)) == (255, 255, 255)
            assert boxed.getpixel((180, 120)) == (255, 0, 0)
            # Cropping removes the large source margins before fitting.
            assert boxed.getpixel((180, 10)) == (255, 0, 0)
            assert boxed.getpixel((160, 120)) == (255, 255, 255)
    register = load_workbook(output)["Register"]
    image = register._images[0]
    assert (image.width, image.height) == (360, 240)
    assert image.anchor._from.col == 1
    assert image.anchor._from.colOff == 4 * 9525
    assert image.anchor._from.rowOff == 3 * 9525
    assert image.anchor.ext.cx == 360 * 9525
    assert image.anchor.ext.cy == 240 * 9525
    assert register.row_dimensions[2].height == 185


def test_workbook_presentation_print_settings_and_chart_placement(tmp_path):
    output = tmp_path / "register.xlsx"
    write_register(
        str(output),
        {"name": "Demo", "model": "Model"},
        [issue(), issue(category="Geometry", severity="critical")],
    )
    workbook = load_workbook(output)
    cover = workbook["Cover"]
    assert cover["A1"].value == "Model review: issue register"
    assert cover["A1"].font.sz == 22
    assert cover["A1"].fill.fgColor.rgb == "001F3864"
    assert cover["A3"].value == "Demo | Model"
    assert cover["A3"].font.sz == 12
    assert "A1:H2" in cover.merged_cells
    assert cover["A17"].value == 1
    assert cover["C17"].value == 1
    assert cover["E17"].value == 0
    assert cover["G17"].value == 0
    assert cover["A20"].value == 2
    assert (
        cover["A23"].value
        == "Snapshots show affected elements in red; other elements are greyed out."
    )
    for sheet in workbook:
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == (1 if sheet.title == "Cover" else 0)
        assert sheet.sheet_properties.pageSetUpPr.fitToPage
        assert sheet.oddFooter.left.text == "Demo"
        assert sheet.oddFooter.right.text == "Page &P of &N"
    for name in ("Cover", "Summary"):
        assert not workbook[name].sheet_view.showGridLines
        assert workbook[name].auto_filter.ref is None
    register = workbook["Register"]
    assert register.print_title_rows == "$1:$1"
    assert register["J2"].value == "Open"
    assert register["J2"].fill.fgColor.rgb == "00FCE4E4"
    assert register["C2"].font.name == "Calibri"
    assert register["C2"].font.sz == 10
    assert register["C2"].alignment.wrap_text
    assert register["C2"].alignment.vertical == "top"
    assert register["C2"].border.left.color.rgb == "00BFBFBF"
    assert register["C2"].fill.fgColor.rgb == "00F2F2F2"
    assert register["C3"].fill.patternType is None
    elements = workbook["Elements"]
    assert elements.freeze_panes == "A2"
    assert elements.auto_filter.ref == "A1:B1"
    summary = workbook["Summary"]
    assert summary["A4"].font.bold
    assert summary["A4"].border.top.style == "thin"
    assert summary["A6"].font.color.rgb == "00FFFFFF"
    chart = summary._charts[0]
    assert chart.anchor._from.col == 8
    assert chart.anchor._from.col >= summary.max_column
    assert chart.legend.position == "b"
    assert chart.gapWidth == 60
    assert chart.anchor.ext.cx == pytest.approx(16 * 360000)
    assert chart.anchor.ext.cy == pytest.approx(9 * 360000)
    assert [series.graphicalProperties.solidFill.srgbClr for series in chart.series] == list(COLORS)
