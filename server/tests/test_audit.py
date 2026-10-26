"""Audit contracts against recorded reads, with no Revit runtime."""

import asyncio
import copy
import json
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from jsonschema import ValidationError
from openpyxl import load_workbook

from revit_model_mcp import server
from revit_model_mcp.audit import to_register_issues
from revit_model_mcp.audit.evaluate import (
    Evaluator,
    LiveReadSource,
    NotAssessed,
    SnapshotReadSource,
    compile_pattern,
)
from revit_model_mcp.audit.report import HEADERS, build_report
from revit_model_mcp.audit.result import build_result, model_identity, model_key
from revit_model_mcp.audit.rules import PACKS, load_rules
from revit_model_mcp.snapshot_report import load_snapshots
from revit_model_mcp.toolsets import TOOLSETS

FIXTURES = Path(__file__).parent / "fixtures" / "snapshots"
MODEL = {
    "title": "Tower",
    "centralPath": r"C:\Private\Tower.rvt",
    "path": None,
    "versionGuid": "family-a",
}


def run(awaitable):
    return asyncio.run(awaitable)


def rule(kind="name_matches", **check):
    return {
        "id": "test",
        "title": "Test rule",
        "severity": "major",
        "check": {
            "type": kind,
            **({"pattern": "^L[0-9]+$"} if kind == "name_matches" else {}),
            **check,
        },
    }


def file(rules=None, **values):
    return {
        "schema": 1,
        "name": "Project checks",
        "rules": rules if rules is not None else [rule()],
        **values,
    }


def read(data, **extra):
    return {"success": True, "data": data, **extra}


def evaluate(r, responses):
    async def reader(tool, arguments):
        response = responses[tool]
        if isinstance(response, Exception):
            raise response
        return response(arguments) if callable(response) else response

    source = AsyncMock()
    source.read.side_effect = reader
    findings = run(Evaluator(source).evaluate([r]))
    return findings[0], source


def test_rule_files_and_all_packs(tmp_path):
    for pack in PACKS:
        loaded = load_rules("pack:" + pack)
        assert loaded["rules"] and loaded["packs"] == [pack]
        assert len(loaded["sha256"]) == 64
    for suffix, content in [
        ("json", json.dumps(file())),
        ("yaml", "schema: 1\nname: Test\nrules: []\n"),
    ]:
        path = tmp_path / ("rules." + suffix)
        path.write_text(content)
        assert load_rules(str(path))["file"] == str(path)
    with pytest.raises(ValueError):
        load_rules("pack:missing")
    with pytest.raises(ValueError):
        load_rules("pack:model-health", file())


@pytest.mark.parametrize(
    "field,value",
    [
        ("schema", 2),
        ("schema", True),
        ("name", ""),
        ("extends", ["missing"]),
        ("references", []),
        ("rules", {}),
        ("unknown", True),
    ],
)
def test_invalid_file_fields(field, value):
    with pytest.raises((ValueError, ValidationError)):
        load_rules(inline=file(**{field: value}))


@pytest.mark.parametrize(
    "field,value",
    [
        ("id", ""),
        ("title", 1),
        ("severity", "fatal"),
        ("scope", {"unknown": "x"}),
        ("check", {"type": "script"}),
        ("threshold", {"min_percent": 101}),
        ("reference", {"unknown": "x"}),
        ("recommendation", False),
        ("enabled", 1),
        ("unknown", "x"),
    ],
)
def test_invalid_rule_fields(field, value):
    r = rule()
    r[field] = value
    with pytest.raises((ValueError, ValidationError)):
        load_rules(inline=file([r]))


@pytest.mark.parametrize(
    "check",
    [
        {"type": "parameter_filled"},
        {"type": "parameter_in", "parameter": {"names": ["Mark"]}},
        {"type": "parameter_matches", "parameter": {"builtin": "MARK"}},
        {"type": "name_matches"},
        {"type": "coordinates"},
        {"type": "sheets"},
        {"type": "worksets"},
        {"type": "model_health"},
        {"type": "parameter_filled", "parameter": {"guid": "bad"}},
        {"type": "parameter_filled", "parameter": {"names": []}},
        {"type": "parameter_filled", "parameter": {"builtin": "X", "names": ["x"]}},
    ],
)
def test_invalid_checks(check):
    r = rule()
    r["check"] = check
    with pytest.raises(ValidationError):
        load_rules(inline=file([r]))


def test_rule_overrides_duplicates_and_limit():
    loaded = load_rules(
        inline=file(
            [
                {
                    "id": "HEALTH-warnings",
                    "severity": "minor",
                    "threshold": {"max_count": 3},
                    "enabled": False,
                }
            ],
            extends=["model-health"],
        )
    )
    inherited = next(r for r in loaded["rules"] if r["id"] == "HEALTH-warnings")
    assert inherited["threshold"] == {"max_count": 3} and inherited["enabled"] is False
    for changed in [
        {"title": "Changed"},
        {"check": {"type": "count"}},
        {"scope": {"level": "L1"}},
        {"reference": {"source": "Changed"}},
    ]:
        with pytest.raises(ValueError, match="Override cannot change"):
            load_rules(
                inline=file([{"id": "HEALTH-warnings", **changed}], extends=["model-health"])
            )
    with pytest.raises(ValueError, match="Duplicate"):
        load_rules(inline=file([rule(), rule()]))
    with pytest.raises(ValueError, match="Incomplete"):
        load_rules(inline=file([{"id": "new"}]))
    rules = [{**rule(), "id": str(i)} for i in range(201)]
    with pytest.raises(ValidationError):
        load_rules(inline=file(rules))
    with pytest.raises(ValueError, match="200"):
        load_rules(inline=file(rules[:200], extends=["model-health"]))


@pytest.mark.parametrize(
    "metric,field",
    [
        ("warnings_total", "warnings"),
        ("in_place_families", "familiesInPlace"),
        ("imported_cad", "cadImports"),
    ],
)
def test_health_metrics(metric, field):
    r = {**rule("model_health", metric=metric), "threshold": {"max_count": 2}}
    f, _ = evaluate(r, {"revit_model_health": read({"counts": {field: 3, "elements": 100}})})
    assert (f["status"], f["measured"], f["population"]) == ("failed", 3, 100)
    f, _ = evaluate(r, {"revit_model_health": read({"counts": {field: None}})})
    assert f["status"] == "not_assessed"


def test_unloaded_links_and_warning_count():
    f, _ = evaluate(
        rule("model_health", metric="unloaded_links"),
        {"revit_links_status": read({"summary": {"rvt": 5, "rvtLoaded": 3}})},
    )
    assert (f["measured"], f["population"]) == (2, 5)
    f, _ = evaluate(
        rule("count", metric="warnings_total"),
        {"revit_list_warnings": read({"total": 4, "groups": []})},
    )
    assert f["status"] == "failed" and f["measured"] == 4


def catalogs():
    return {
        "revit_list_catalog": lambda args: read(
            {
                "items": [{"name": "Стены", "builtin": "OST_Walls"}]
                if args["section"] == "categories"
                else [
                    {
                        "name": "Марка",
                        "builtin": "ALL_MODEL_MARK",
                        "guid": "00000000-0000-0000-0000-000000000001",
                    }
                ]
            }
        )
    }


def fill_rule():
    return {
        **rule("parameter_filled", parameter={"builtin": "ALL_MODEL_MARK"}),
        "scope": {"categories": [{"builtin": "OST_Walls"}]},
        "threshold": {"min_percent": 100},
    }


def test_live_fill_shape_and_resolution():
    r = fill_rule()
    f, source = evaluate(
        r,
        {
            **catalogs(),
            "revit_parameter_fill_check": read(
                {
                    "parameters": [
                        {
                            "name": "Марка",
                            "elements": 10,
                            "filled": 8,
                            "empty": 1,
                            "missing": 1,
                            "emptySampleIds": [12],
                            "missingSampleIds": [13],
                        }
                    ]
                }
            ),
        },
    )
    assert (f["measured"], f["population"], f["element_ids"]) == (80, 10, [12, 13])
    assert not f["element_ids_truncated"]
    assert source.read.call_args.args == (
        "revit_parameter_fill_check",
        {"categories": ["Стены"], "parameters": ["Марка"]},
    )
    for reference in (
        {"guid": "00000000-0000-0000-0000-000000000001"},
        {"names": ["Missing", "Марка"]},
    ):
        r["check"]["parameter"] = reference
        f, _ = evaluate(
            r,
            {
                **catalogs(),
                "revit_parameter_fill_check": read({"rows": [{"total": 1, "filled": 1}]}),
            },
        )
        assert f["status"] == "passed"


@pytest.mark.parametrize(
    "kind,check",
    [
        ("name_matches", {}),
        ("parameter_in", {"parameter": {"names": ["Марка"]}, "values": ["L01"]}),
        ("parameter_matches", {"parameter": {"names": ["Марка"]}, "pattern": "^L[0-9]+$"}),
        ("count", {}),
    ],
)
def test_element_checks(kind, check):
    field = "name" if kind in ("name_matches", "count") else "Марка"
    rows = [
        {"id": i, "values": {field: {"hasValue": True, "value": value}}}
        for i, value in [(1, "L01"), (2, "Bad")]
    ]
    f, _ = evaluate(
        rule(kind, **check),
        {**catalogs(), "revit_query_elements": read({"elements": rows, "hasMore": False})},
    )
    assert f["status"] == "failed"
    assert f["measured"] == (2 if kind == "count" else 1)
    assert f["element_ids"] == ([] if kind == "count" else [2])


def test_query_paging_cache_skips_and_budget():
    source = AsyncMock()
    source.read.return_value = read(
        {
            "elements": [{"id": i, "values": {"name": {"value": "Bad"}}} for i in range(501)],
            "hasMore": True,
        },
        skipped=[{"what": "x", "reason": "missing"}],
        skippedCount=3,
    )
    evaluator = Evaluator(source)
    f, g = run(evaluator.evaluate([rule(), {**rule(), "id": "second"}]))
    assert source.read.await_count == 20
    assert f["status"] == g["status"] == "partial"
    assert len(f["element_ids"]) == 500 and f["element_ids_truncated"]
    assert f["skippedCount"] == 60 and f["skipped"]
    assert source.read.call_args.args[1]["offset"] == 9500
    source.reset_mock()
    f = run(Evaluator(source, budget_seconds=0).evaluate([rule()]))[0]
    assert f["status"] == "partial" and source.read.await_count == 0


def test_read_timeout_and_unknown_command():
    f, _ = evaluate(rule(), {"revit_query_elements": TimeoutError("budget")})
    assert f["status"] == "partial"
    f, _ = evaluate(
        rule("sheets", metric="missing_title_block"),
        {"revit_list_sheets": RuntimeError("unknown command")},
    )
    assert f["status"] == "not_assessed" and "unknown command" in f["reason"]


@pytest.mark.parametrize("pattern", ["a" * 201, "(a+)+$", "(a{1,3})*", "["])
def test_regex_guard(pattern):
    with pytest.raises(NotAssessed):
        compile_pattern(pattern)


def test_regex_timeout_is_partial():
    matcher = AsyncMock()  # A deterministic two-second guard failure without slowing the suite.
    matcher.fullmatch = lambda *a, **kw: (_ for _ in ()).throw(TimeoutError())
    with patch("revit_model_mcp.audit.evaluate.compile_pattern", return_value=matcher):
        f, _ = evaluate(
            rule(),
            {
                "revit_query_elements": read(
                    {"elements": [{"id": 1, "values": {"name": "x"}}], "hasMore": False}
                )
            },
        )
    assert f["status"] == "partial"


@pytest.mark.parametrize(
    "metric,expected",
    [("missing_title_block", [1]), ("unplaced_views", [3]), ("name_matches", [2])],
)
def test_sheet_checks(metric, expected):
    r = rule("sheets", metric=metric, pattern="^A.*$", field="number")
    f, source = evaluate(
        r,
        {
            "revit_list_sheets": read(
                {
                    "sheets": [
                        {"id": 1, "number": "A01", "name": "One", "titleBlock": None},
                        {"id": 2, "number": "B01", "name": "Two", "titleBlock": "A1"},
                    ],
                    "unplacedViews": [{"id": 3, "name": "Plan", "viewType": "FloorPlan"}],
                    "hasMore": False,
                }
            )
        },
    )
    assert f["status"] == "failed" and f["element_ids"] == expected
    assert source.read.call_args.args[1]["include_unplaced_views"] is True


@pytest.mark.parametrize("has_issues", [False, True])
def test_sheets_and_views_pack_evaluates_recorded_responses(has_issues):
    responses = {
        "revit_list_sheets": read(
            {
                "sheets": [
                    {
                        "id": 1,
                        "number": "A01",
                        "name": "Ground floor",
                        "titleBlock": None if has_issues else "A1",
                        "revisionCount": 2,
                        "viewportCount": 1,
                        "viewIds": [2],
                    }
                ],
                "unplacedViews": [{"id": 3, "name": "Doors", "viewType": "Schedule"}]
                if has_issues
                else [],
                "total": 1,
                "offset": 0,
                "limit": 200,
                "hasMore": False,
            },
            skipped=[],
            skippedCount=0,
        ),
        "revit_list_views": read(
            {
                "views": [
                    {
                        "id": 2,
                        "name": "Unnamed" if has_issues else "Ground floor",
                        "type": "FloorPlan",
                    }
                ],
                "total": 1,
                "processed": 1,
            }
        ),
    }
    findings = []
    for pack_rule in load_rules("pack:sheets-and-views")["rules"]:
        finding, _ = evaluate(pack_rule, responses)
        findings.append(finding)
    assert [finding["status"] for finding in findings] == ["failed" if has_issues else "passed"] * 3
    assert [finding["element_ids"] for finding in findings] == (
        [[1], [3], [2]] if has_issues else [[], [], []]
    )


def test_sheet_checks_follow_pages_and_keep_first_page_unplaced_views():
    def page(arguments):
        offset = arguments["offset"]
        return read(
            {
                "sheets": [
                    {
                        "id": view_id,
                        "number": f"A{view_id}",
                        "name": "Plan",
                        "titleBlock": "A1",
                        "revisionCount": 0,
                        "viewportCount": 0,
                        "viewIds": [],
                    }
                    for view_id in range(offset + 1, min(offset + 201, 202))
                ],
                "unplacedViews": [{"id": 3, "name": "Doors", "viewType": "Schedule"}],
                "total": 201,
                "offset": offset,
                "limit": 200,
                "hasMore": offset == 0,
            }
        )

    finding, source = evaluate(rule("sheets", metric="unplaced_views"), {"revit_list_sheets": page})
    assert finding["status"] == "failed"
    assert finding["element_ids"] == [3]
    assert source.read.await_count == 2
    assert [call.args[1]["offset"] for call in source.read.await_args_list] == [0, 200]


def view_name_rule():
    return next(r for r in load_rules("pack:sheets-and-views")["rules"] if r["id"] == "VIEWS-names")


@pytest.mark.parametrize(
    "names,expected_ids",
    [
        (["Floor Plan Copy 1", "Unnamed", "{3D}"], [101, 102, 103]),
        (["Ground floor", "North elevation", "Coordination model"], []),
    ],
)
def test_view_names_use_view_reader(names, expected_ids):
    views = [
        {"id": i, "name": name, "type": view_type, "isTemplate": False, "template": "Architectural"}
        for i, name, view_type in zip(
            [101, 102, 103], names, ["FloorPlan", "Elevation", "ThreeD"], strict=True
        )
    ]
    f, source = evaluate(
        view_name_rule(),
        {"revit_list_views": read({"views": views, "total": 3, "processed": 3})},
    )
    assert f["status"] == ("failed" if expected_ids else "passed")
    assert (f["measured"], f["population"], f["element_ids"]) == (
        len(expected_ids),
        3,
        expected_ids,
    )
    assert f["reads"] == ["revit_list_views"]
    source.read.assert_awaited_once_with(
        "revit_list_views", {"include_templates": False, "offset": 0, "limit": 500}
    )


@pytest.mark.parametrize("include_templates", [False, True])
def test_view_templates_are_excluded_unless_requested(include_templates):
    r = view_name_rule()
    r["scope"]["include_templates"] = include_templates
    load_rules(inline=file([{k: v for k, v in r.items() if k != "pack"}]))
    f, _ = evaluate(
        r,
        {
            "revit_list_views": read(
                {
                    "views": [
                        {
                            "id": 10,
                            "name": "Ground floor",
                            "type": "FloorPlan",
                            "isTemplate": False,
                        },
                        {"id": 11, "name": "Unnamed", "type": "FloorPlan", "isTemplate": True},
                    ]
                }
            )
        },
    )
    assert f["status"] == ("failed" if include_templates else "passed")
    assert f["population"] == (2 if include_templates else 1)
    assert f["element_ids"] == ([11] if include_templates else [])


def test_view_names_page_through_all_views():
    def page(arguments):
        offset = arguments["offset"]
        return read(
            {
                "views": [
                    {
                        "id": offset + 1,
                        "name": "Unnamed" if offset == 21 else "Ground floor",
                        "type": "FloorPlan",
                        "isTemplate": False,
                    }
                ],
                "hasMore": offset < 21,
            }
        )

    f, source = evaluate(view_name_rule(), {"revit_list_views": page})
    assert f["status"] == "failed"
    assert (f["population"], f["element_ids"]) == (22, [22])
    assert [call.args[1]["offset"] for call in source.read.call_args_list] == list(range(22))


@pytest.mark.parametrize(
    "response",
    [RuntimeError("unknown command"), {"success": False, "error": "unknown command"}],
)
def test_view_names_unavailable_reader(response):
    f, _ = evaluate(view_name_rule(), {"revit_list_views": response})
    assert f["status"] == "not_assessed"
    assert "unknown command" in f["reason"]
    assert f["reads"] == ["revit_list_views"]


def test_view_names_partial_and_empty_reads():
    for response in (
        {
            "success": False,
            "partial": True,
            "data": {"views": [{"id": 1, "name": "Ground floor", "type": "FloorPlan"}]},
        },
        read({"views": [], "hasMore": True}),
    ):
        f, _ = evaluate(view_name_rule(), {"revit_list_views": response})
        assert f["status"] == "partial"
    f, _ = evaluate(view_name_rule(), {"revit_list_views": read({"views": []})})
    assert f["status"] == "passed" and f["population"] == 0


def test_live_view_names_reader():
    channel = AsyncMock()
    channel.execute.return_value = read(
        {
            "views": [{"id": 7, "name": "{3D}", "type": "ThreeD", "template": None}],
            "total": 1,
            "processed": 1,
        }
    )
    with patch.object(server, "channel", channel):
        f = run(Evaluator(LiveReadSource("Tower", 41)).evaluate([view_name_rule()]))[0]
    assert f["status"] == "failed" and f["element_ids"] == [7]
    job, timeout = channel.execute.call_args.args
    assert job.command == "list-views" and timeout == 41
    assert job.payload == {
        "command": "list-views",
        "offset": 0,
        "limit": 500,
        "targetDocument": "Tower",
    }
    with pytest.raises(NotAssessed, match="cannot include view templates"):
        run(LiveReadSource().read("revit_list_views", {"include_templates": True}))


def test_coordinates_and_worksets():
    r = {**rule("coordinates", metric="link_offset_mm"), "threshold": {"max_count": 4}}
    f, _ = evaluate(
        r,
        {
            "revit_shared_coordinates": read(
                {
                    "sharedSiteFromLinks": [{"offsetMm": {"x": 3, "y": 4, "z": 0}}],
                    "linkInstancesTotal": 1,
                }
            )
        },
    )
    assert (f["status"], f["measured"], f["population"]) == ("failed", 5, 1)
    r["threshold"]["max_count"] = 5
    f, _ = evaluate(
        r, {"revit_shared_coordinates": read({"sharedSiteFromLinks": [], "linkInstancesTotal": 2})}
    )
    assert f["status"] == "partial"
    for r in [
        rule("coordinates", metric="present"),
        rule("worksets", metric="closed"),
        {**rule(), "enabled": False},
    ]:
        f, source = evaluate(r, {})
        assert f["status"] == "not_assessed" and source.read.await_count == 0


def test_unresolved_references_and_unsupported_scope():
    f, _ = evaluate(fill_rule(), {"revit_list_catalog": read({"items": []})})
    assert f["status"] == "not_assessed"
    r = {**rule("model_health", metric="warnings_total"), "scope": {"level": "L1"}}
    f, source = evaluate(r, {})
    assert f["status"] == "not_assessed" and source.read.await_count == 0


def test_snapshot_v1_fixture():
    snapshot = load_snapshots(str(FIXTURES / "current"))[0][1]
    source = SnapshotReadSource(snapshot)
    loaded = load_rules("pack:model-health")
    findings = run(Evaluator(source).evaluate(loaded["rules"]))
    by_id = {f["id"]: f for f in findings}
    assert by_id["HEALTH-warnings"]["measured"] == snapshot["warnings"]["total"]
    assert by_id["HEALTH-in-place"]["measured"] == snapshot["families"]["inPlace"]
    assert by_id["HEALTH-links"]["status"] == "not_assessed"
    r = fill_rule()
    r["scope"] = {"categories": [{"names": ["Doors"]}]}
    r["check"]["parameter"] = {"names": ["Mark"]}
    f = run(Evaluator(source).evaluate([r]))[0]
    assert f["measured"] == 80 and f["element_ids"] == [1, 2]
    for tool in ("revit_query_elements", "revit_links_status", "revit_list_sheets"):
        with pytest.raises(NotAssessed):
            run(source.read(tool, {}))
    r["scope"]["level"] = "L1"
    assert run(Evaluator(source).evaluate([r]))[0]["status"] == "not_assessed"
    with pytest.raises(ValueError):
        SnapshotReadSource({"schemaVersion": 2})


def finding(status="failed", ids=None):
    return {
        **rule(),
        "pack": "model-health",
        "status": status,
        "measured": 2,
        "population": 5,
        "threshold": {"max_count": 0},
        "element_ids": [1, 2] if ids is None else ids,
        "element_ids_truncated": False,
        "reads": ["revit_query_elements"],
        "skipped": [],
        "skippedCount": 0,
    }


def test_history_and_deltas(tmp_path, monkeypatch):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path))
    rules = load_rules(inline=file())
    first = build_result(MODEL, rules, [finding()], evaluated_at="2026-01-01T00:00:00Z")
    second = build_result(MODEL, rules, [finding(ids=[2, 3])], evaluated_at="2026-01-02T00:00:00Z")
    assert {d["element_id"]: d["status"] for d in second["deltas"]} == {
        "1": "fixed",
        "2": "unchanged",
        "3": "new",
    }
    assert Path(first["savedTo"]).is_file()
    with pytest.raises(ValueError, match="already exists"):
        build_result(MODEL, rules, [finding()], evaluated_at="2026-01-01T00:00:00Z")
    changed = {**rules, "sha256": "changed"}
    result = build_result(
        MODEL, changed, [finding("passed", [])], compare_with=second["savedTo"], save=False
    )
    assert result["summary"]["rulesChanged"] and result["deltas"]
    assert result["savedTo"] is None
    different = {**MODEL, "versionGuid": "another"}
    result = build_result(different, rules, [finding()], compare_with=first["savedTo"], save=False)
    assert "deltas" not in result and "different" in result["summary"]["comparison"]
    assert model_key(MODEL) != model_key({**MODEL, "centralPath": "other"})


@pytest.mark.parametrize(
    "before,after", [("failed", "not_assessed"), ("not_assessed", "passed"), ("failed", "partial")]
)
def test_incomplete_never_fixed(tmp_path, monkeypatch, before, after):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path))
    rules = load_rules(inline=file())
    first = build_result(MODEL, rules, [finding(before)], compare_with="none")
    result = build_result(
        MODEL, rules, [finding(after, [])], compare_with=first["savedTo"], save=False
    )
    assert not any(d["status"] == "fixed" for d in result["deltas"])


def test_model_level_delta_and_history_redaction(tmp_path, monkeypatch):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path))
    monkeypatch.setenv("REVIT_MCP_REDACT_PATHS", "1")
    rules = load_rules(inline=file())
    first = build_result(MODEL, rules, [finding(ids=[])], compare_with="none")
    assert first["model"]["centralPath"] == "Tower.rvt"
    assert Path(first["savedTo"]).name == first["savedTo"]
    stored = next(tmp_path.rglob("*.json"))
    assert "Private" not in stored.read_text()
    second = build_result(MODEL, rules, [finding("passed", [])], save=False)
    assert second["deltas"][0]["status"] == "fixed"
    assert second["deltas"][0]["element_id"] == ""
    assert (
        model_identity({"passport": {"title": "X"}, "source": {"path": "a.rvt"}})["path"] == "a.rvt"
    )


def test_reports_and_register_mapping(tmp_path, monkeypatch):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path / "history"))
    f = finding()
    f.update(
        title="=Formula",
        recommendation='=HYPERLINK("bad")',
        reference={"source": "BEP", "clause": "4.2", "text": "<Fill Mark>"},
    )
    rules = load_rules(inline=file())
    result = build_result(MODEL, rules, [f], compare_with="none")
    issues = to_register_issues(result)
    assert issues[0]["title"] == "=Formula" and issues[0]["category"] == "model-health"
    assert issues[0]["requirement_source"] == "BEP"
    assert issues[0]["requirement"] == "4.2: <Fill Mark>"
    assert to_register_issues({**result, "findings": [finding("not_assessed")]}) == []
    outputs = build_report(result["savedTo"], str(tmp_path / "audit"), "both")["files"]
    workbook = load_workbook(outputs[0])
    assert workbook.sheetnames == ["Summary", "Findings", "Rules", "Skipped"]
    assert [c.value for c in workbook["Findings"][1]] == HEADERS
    assert workbook["Findings"]["H2"].value == f["recommendation"]
    assert workbook["Findings"]["H2"].data_type == "s"
    document = Path(outputs[1]).read_text()
    assert "&lt;Fill Mark&gt;" in document
    assert 'src="http' not in document and 'href="http' not in document
    with pytest.raises(ValueError, match="already exists"):
        build_report(result["savedTo"], str(tmp_path / "audit"))
    result["deltas"] = [{"rule": "test", "element_id": "1", "status": "new", "severity": "major"}]
    path = tmp_path / "with-deltas.json"
    path.write_text(json.dumps(result))
    outputs = build_report(str(path), str(tmp_path / "second.xlsx"), "xlsx")["files"]
    assert "Deltas" in load_workbook(outputs[0]).sheetnames


def test_tool_membership():
    for tool, group in [("revit_model_audit", "read"), ("revit_audit_report", "batch")]:
        assert [g for g, members in TOOLSETS.items() if tool in members] == [group]


def test_live_adapter_and_audit_addressing(tmp_path, monkeypatch):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path))
    channel = AsyncMock()
    snapshot = {
        "schemaVersion": 1,
        "passport": {
            "title": "Tower",
            "centralPath": MODEL["centralPath"],
            "versionGuid": "family-a",
        },
        "source": {"path": "Tower.rvt"},
    }
    channel.execute.side_effect = [
        read(snapshot),
        read({"counts": {"warnings": 3, "elements": 10}}),
    ]
    with patch.object(server, "channel", channel):
        result = run(
            server.revit_model_audit(
                rules_inline=file([rule("model_health", metric="warnings_total")]),
                compare_with="none",
                save=False,
                document="Tower",
                process_id=123,
            )
        )
    assert result["findings"][0]["status"] == "failed"
    for call in channel.execute.call_args_list:
        assert call.args[0].payload["targetDocument"] == "Tower"
        assert call.args[0].payload["targetProcessId"] == 123
    channel.execute.reset_mock()
    with patch.object(server, "channel", channel), pytest.raises(Exception):
        run(server.revit_model_audit(rules_inline=file([rule("unknown")])))
    channel.execute.assert_not_awaited()
    channel.execute.side_effect = None
    channel.execute.return_value = {"success": False, "error": "unknown command"}
    with patch.object(server, "channel", channel):
        f = run(
            Evaluator(LiveReadSource("Tower")).evaluate(
                [rule("sheets", metric="missing_title_block")]
            )
        )[0]
    assert f["status"] == "not_assessed"
    assert channel.execute.call_args.args[0].payload["includeUnplacedViews"] is True


def test_default_state_root(monkeypatch):
    from revit_model_mcp.audit.result import audit_directory
    from revit_model_mcp.updates import state_path

    monkeypatch.delenv("REVIT_MCP_AUDIT_DIR", raising=False)
    assert audit_directory() == state_path().parent / "audits"


def test_history_does_not_mutate_inputs(tmp_path, monkeypatch):
    monkeypatch.setenv("REVIT_MCP_AUDIT_DIR", str(tmp_path))
    model, rules = copy.deepcopy(MODEL), load_rules(inline=file())
    build_result(model, rules, [finding()], save=False, compare_with="none")
    assert model == MODEL and rules["file"] == "inline"


def test_regex_nested_groups_and_literal_quantifiers():
    for pattern in ("((a)+)+$", "(a?)*", "(?:a{1,2})+"):
        with pytest.raises(NotAssessed):
            compile_pattern(pattern)
    assert compile_pattern(r"^[+*?]+$").fullmatch("+*?")
    assert compile_pattern(r"^L\d{2} - .+$").fullmatch("L01 - Ground")


def test_regex_matches_receive_guard_and_run_deadline():
    evaluator = Evaluator(AsyncMock())
    matcher = AsyncMock()
    matcher.fullmatch = lambda value, **kw: kw["timeout"]
    assert evaluator.matches(matcher, "value") == 2
    evaluator.deadline = 0
    with pytest.raises(TimeoutError):
        evaluator.matches(matcher, "value")


def test_partial_read_cannot_pass():
    f, _ = evaluate(
        rule("model_health", metric="warnings_total"),
        {"revit_model_health": read({"counts": {"warnings": 0, "elements": 10}}, partial=True)},
    )
    assert f["status"] == "partial" and f["skipped"]


def test_live_fill_reader_records_payload():
    channel = AsyncMock()
    channel.execute.return_value = read({"parameters": []})
    with patch.object(server, "channel", channel):
        run(
            LiveReadSource("Tower", 41).read(
                "revit_parameter_fill_check", {"categories": ["Walls"], "parameters": ["Mark"]}
            )
        )
    job, timeout = channel.execute.call_args.args
    assert job.command == "parameter-fill-check"
    assert job.payload["parameters"] == ["Mark"] and job.payload["categories"] == ["Walls"]
    assert timeout == 41
    assert job.payload["targetDocument"] == "Tower"


def test_snapshot_missing_parameter_and_skipped_evidence():
    snapshot = load_snapshots(str(FIXTURES / "current"))[0][1]
    snapshot["skipped"] = [{"what": "counts.cadImports", "reason": "unavailable"}]
    snapshot["skippedCount"] = 5
    source = SnapshotReadSource(snapshot)
    f = run(Evaluator(source).evaluate([rule("model_health", metric="warnings_total")]))[0]
    assert f["status"] == "partial" and f["skippedCount"] == 5
    with pytest.raises(NotAssessed):
        run(
            source.read(
                "revit_parameter_fill_check", {"categories": ["Doors"], "parameters": ["Missing"]}
            )
        )


def test_rules_hash_changes_with_threshold():
    first = load_rules(inline=file())
    second_rule = {**rule(), "threshold": {"max_count": 10}}
    second = load_rules(inline=file([second_rule]))
    assert first["sha256"] != second["sha256"]


def test_parameter_fill_requires_percentage_threshold():
    r = fill_rule()
    r["threshold"] = {"max_count": 1}
    with pytest.raises(ValidationError):
        load_rules(inline=file([r]))


def test_inherited_fill_threshold_is_validated_after_merge():
    with pytest.raises(ValidationError):
        load_rules(
            inline=file(
                [{"id": "INFO-required", "threshold": {"max_count": 1}}],
                extends=["iso19650-information"],
            )
        )
