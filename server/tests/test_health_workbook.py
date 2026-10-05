import asyncio
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook
from PIL import Image

from revit_model_mcp import server
from revit_model_mcp.health_workbook import CHECKS, write_health_workbook


@pytest.fixture
def health():
    return {
        "success": True,
        "data": {
            "counts": dict.fromkeys(
                (
                    "elements",
                    "warnings",
                    "warningGroups",
                    "levels",
                    "grids",
                    "views",
                    "viewsNotOnSheets",
                    "viewTemplates",
                    "sheets",
                    "rooms",
                    "roomsUnplaced",
                    "roomsNotEnclosed",
                    "families",
                    "familiesInPlace",
                    "familyTypesUnused",
                    "groupsModel",
                    "groupsDetail",
                    "groupTypes",
                    "designOptions",
                    "worksets",
                    "linksRvt",
                    "linksCad",
                    "cadImports",
                    "images",
                ),
                0,
            ),
            "fileName": "Demo.rvt",
            "fileSizeBytes": 10485760,
            "isWorkshared": True,
            "projectInfo": {
                "name": "Demo",
                "number": "01",
                "client": "Client",
                "address": "Address",
                "buildingName": "Building",
                "status": "Review",
                "author": "Author",
            },
            "revitBuild": "20260101",
            "revitVersion": "2026",
            "topWarnings": [{"count": 3, "text": "Overlapping walls"}],
        },
        "skipped": [],
        "skippedCount": 0,
    }


def groups():
    return [
        {
            "count": i,
            "text": f"Warning {i}",
            "severity": "Warning",
            "affectedElementCount": 60,
            "elements": [{"id": j} for j in range(1, 61)],
        }
        for i in range(1, 8)
    ]


def test_workbook_checks_counts_and_presentation(tmp_path, health):
    health["data"]["counts"].update(
        views=100, viewsNotOnSheets=20, familyTypesUnused=300, images=20, groupsModel=50
    )
    result = write_health_workbook(str(tmp_path / "health.xlsx"), health, groups())
    book = load_workbook(result["path"])
    assert book.sheetnames == ["Health", "Warnings", "Counts"]
    sheet = book["Health"]
    assert sheet["A1"].value == "Model health check"
    assert "Demo.rvt | Revit 2026 | 10.0 MB" in sheet["A3"].value
    assert sheet["B6"].value == "Demo"
    assert sheet["B10"].value == "Author"
    assert sheet["B27"].value == "10 of 10 checks pass"
    for row in range(17, 27):
        assert sheet.cell(row, 4).value == "Pass"
        assert sheet.cell(row, 4).fill.fgColor.rgb == "00C6EFCE"
        assert sheet.cell(row, 4).font.color.rgb == "00006100"
    assert sheet["B18"].value == 0.2
    assert sheet["C18"].number_format == "0%"
    assert sheet.freeze_panes == "A17"
    assert sheet.print_title_rows == "$16:$16"
    counts = dict(book["Counts"].values)
    assert counts["views"] == 100
    assert len(counts) == len(health["data"]["counts"]) + 1
    warnings = book["Warnings"]
    assert [cell.value for cell in warnings[1]][:4] == [
        "Count",
        "Warning text",
        "Element IDs",
        "Snapshot",
    ]
    assert warnings["A8"].value == 7
    assert warnings["C2"].value == ", ".join(map(str, range(1, 31))) + " +30 more"
    chart = warnings._charts[0]
    assert chart.type == "bar"
    assert chart.anchor._from.col == 8
    assert chart.series[0].val.numRef.f == "'Warnings'!$G$2:$G$8"
    assert warnings["G2"].value == 7
    assert warnings.freeze_panes == "A2"
    for tab in book:
        assert tab.page_setup.orientation == "landscape"
        assert tab.page_setup.fitToWidth == 1
        assert tab.sheet_properties.pageSetUpPr.fitToPage
        assert tab.oddFooter.right.text == "Page &P of &N"


@pytest.mark.parametrize("unavailable", [False, True])
def test_review_and_unavailable_counts(tmp_path, health, unavailable):
    for _, key, threshold, _ in CHECKS:
        health["data"]["counts"][key] = None if unavailable else threshold + 1
    health["data"]["counts"]["views"] = 1
    result = write_health_workbook(str(tmp_path / "health.xlsx"), health, [])
    sheet = load_workbook(result["path"])["Health"]
    assert sheet["B27"].value == "0 of 10 checks pass"
    for row in range(17, 27):
        assert sheet.cell(row, 4).value == "Review"
        assert sheet.cell(row, 4).fill.fgColor.rgb == "00FFEB9C"
        assert sheet.cell(row, 4).font.color.rgb == "009C6500"
        if unavailable:
            assert sheet.cell(row, 2).value == "Unavailable"


def test_snapshots_and_failures(tmp_path, health):
    png = tmp_path / "snapshot.png"
    Image.new("RGB", (20, 40), "red").save(png)
    result = write_health_workbook(
        str(tmp_path / "health.xlsx"), health, groups(), {6: png, 5: tmp_path / "missing.png"}
    )
    sheet = load_workbook(result["path"])["Warnings"]
    assert result["snapshotCount"] == 1
    assert len(result["warnings"]) == 1
    assert len(sheet._images) == 1
    image = sheet._images[0]
    assert (image.width, image.height) == (360, 240)
    assert image.anchor._from.col == 3
    assert sheet.row_dimensions[8].height == 185
    assert "Snapshot unavailable" in sheet["D7"].value


@pytest.mark.parametrize("destination", ["existing.xlsx", "missing/report.xlsx", "report.csv"])
def test_destination_validation_before_revit(tmp_path, destination, health):
    path = tmp_path / destination
    (tmp_path / "existing.xlsx").write_bytes(b"original")
    with patch.object(server.channel, "execute", new=AsyncMock()) as execute:
        with pytest.raises(server.ToolError):
            asyncio.run(server.revit_model_health(save_to=str(path)))
    execute.assert_not_awaited()
    with pytest.raises(ValueError):
        write_health_workbook(str(path), health, [])
    assert (tmp_path / "existing.xlsx").read_bytes() == b"original"


def test_response_without_save_to_is_unchanged(health):
    with patch.object(server.channel, "execute", new=AsyncMock(return_value=health)) as execute:
        result = asyncio.run(server.revit_model_health())
    assert result == health
    assert execute.await_count == 1
    assert execute.await_args.args[0].command == "model-health"


def test_tool_alias_top_five_captures_and_target(tmp_path, health):
    calls = []

    async def execute(job, *_):
        calls.append(job)
        if job.command == "model-health":
            return health
        if job.command == "list-warnings":
            return {"data": {"groups": groups()}}
        assert job.command == "capture-elements"
        Image.new("RGB", (20, 20), "red").save(job.save_to)
        return {"data": {"localPath": job.save_to}}

    with patch.object(server.channel, "execute", side_effect=execute):
        result = asyncio.run(
            server.mcp.call_tool(
                "revit_model_health",
                {
                    "saveTo": str(tmp_path / "health.xlsx"),
                    "document": "Demo",
                    "process_id": 42,
                },
            )
        )
    assert not result.is_error
    assert [job.command for job in calls] == [
        "model-health",
        "list-warnings",
        *(["capture-elements"] * 5),
    ]
    assert calls[1].payload["includeElements"] is True
    for job in calls:
        assert job.payload["targetDocument"] == "Demo"
        assert job.payload["targetProcessId"] == 42
    assert [Path(job.save_to).stem for job in calls[2:]] == [
        f"snapshot_{i}" for i in (6, 5, 4, 3, 2)
    ]
    for job in calls[2:]:
        assert job.payload["elementIds"] == list(range(1, 51))
        assert job.payload["mode"] == "3d"
        assert not Path(job.save_to).exists()
    book = load_workbook(tmp_path / "health.xlsx")
    assert len(book["Warnings"]._images) == 5


def test_capture_budget_and_errors_keep_health_response(tmp_path, health):
    captures = []

    async def execute(job, *_):
        if job.command == "model-health":
            return health
        if job.command == "list-warnings":
            return {"data": {"groups": groups()}}
        captures.append(job)
        await asyncio.sleep(1)

    with (
        patch.object(server.channel, "execute", side_effect=execute),
        patch.object(server, "HEALTH_CAPTURE_BUDGET_SECONDS", 0.01),
    ):
        result = asyncio.run(server.revit_model_health(save_to=str(tmp_path / "health.xlsx")))
    assert len(captures) == 1
    assert result["data"] == health["data"]
    assert result["workbook"]["snapshotCount"] == 0
    assert len(result["workbook"]["warnings"]) == 5
    assert "time budget exhausted" in result["workbook"]["warnings"][-1]


def test_capture_failure_is_a_warning(tmp_path, health):
    async def execute(job, *_):
        if job.command == "model-health":
            return health
        if job.command == "list-warnings":
            return {"data": {"groups": groups()[:1]}}
        raise server.ToolError("capture failed")

    with patch.object(server.channel, "execute", side_effect=execute):
        result = asyncio.run(server.revit_model_health(save_to=str(tmp_path / "health.xlsx")))
    assert result["workbook"]["snapshotCount"] == 0
    assert result["workbook"]["warnings"] == [
        "Warning group 1: Snapshot unavailable: capture failed"
    ]


def test_chart_is_limited_to_top_ten_and_text_is_literal(tmp_path, health):
    health["data"]["projectInfo"]["name"] = "=project"
    warnings = [
        {"count": count, "text": f"=Warning {count}", "elements": []} for count in range(1, 13)
    ]
    result = write_health_workbook(str(tmp_path / "health.xlsx"), health, warnings)
    book = load_workbook(result["path"])
    sheet = book["Warnings"]
    assert sheet["A13"].value == 12
    assert sheet["G2"].value == 12
    assert sheet["G11"].value == 3
    assert sheet._charts[0].series[0].val.numRef.f == "'Warnings'!$G$2:$G$11"
    for cell in (book["Health"]["B6"], sheet["B2"], sheet["F2"]):
        assert cell.value.startswith("=")
        assert cell.data_type == "s"
