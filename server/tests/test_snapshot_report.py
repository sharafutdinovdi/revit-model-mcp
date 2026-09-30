import asyncio
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import load_workbook

from revit_model_mcp import server
from revit_model_mcp.snapshot_report import HEADERS, build_report

FIXTURES = Path(__file__).parent / "fixtures" / "snapshots"
CURRENT = FIXTURES / "current"
PREVIOUS = FIXTURES / "previous"


@pytest.mark.parametrize("with_previous", [False, True])
def test_workbook_sheets_and_values(tmp_path, with_previous):
    output = tmp_path / "report.xlsx"
    result = build_report(
        str(CURRENT),
        str(output),
        str(PREVIOUS) if with_previous else None,
        [
            {
                "model": "=Model A",
                "severity": "High",
                "rule": "=Rule",
                "element_ids": [3, 4],
                "recommendation": "=Fix",
            }
        ],
    )
    workbook = load_workbook(output, data_only=False)
    expected = ["Summary", "Warnings", "Families", "Parameters", "Skipped"]
    if with_previous:
        expected.append("Changes")
    expected.append("Findings")
    assert workbook.sheetnames == expected
    assert result == {"outputPath": str(output), "snapshotCount": 2, "sheets": expected}
    for sheet in workbook:
        assert [cell.value for cell in sheet[1]] == HEADERS[sheet.title]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == sheet.dimensions
        assert all(sheet.column_dimensions[cell.column_letter].width >= 12 for cell in sheet[1])
        assert all(sheet.column_dimensions[cell.column_letter].width <= 55 for cell in sheet[1])
        assert all(cell.font.bold for cell in sheet[1])
    assert {name: workbook[name].max_row for name in expected} == {
        **{
            "Summary": 3,
            "Warnings": 3,
            "Families": 5,
            "Parameters": 7,
            "Skipped": 2,
            "Findings": 2,
        },
        **({"Changes": 3} if with_previous else {}),
    }
    summary = workbook["Summary"]
    assert [summary.cell(2, col).value for col in (1, 2, 3, 4, 6, 7, 8, 16, 19, 20)] == [
        "Model A",
        2024,
        2026,
        True,
        14,
        "2026-09-29T10:00:00Z",
        None,
        10,
        2,
        1,
    ]
    assert summary["H3"].value == "2026-09-30T11:00:00Z"
    assert summary["I3"].value == "RS User"
    warnings = workbook["Warnings"]
    assert warnings["D2"].value == "1, 2 [truncated]"
    assert warnings["B2"].value == "=Check constraints"
    assert warnings["B2"].data_type != "f"
    families = workbook["Families"]
    assert families["I2"].value == "in-warnings, large"
    assert families["B2"].data_type != "f"
    assert families["B2"].value == "=Door"
    assert families["F2"].value == 2
    parameters = workbook["Parameters"]
    assert parameters["H2"].value == 80
    assert parameters["H4"].value == 0
    assert parameters["I2"].value == "1, 2"
    assert workbook["Skipped"]["B2"].data_type != "f"
    assert workbook["Skipped"]["B2"].value == "=Hidden set"
    findings = workbook["Findings"]
    assert findings["D2"].value == "3, 4"
    assert findings["C2"].value == "=Rule"
    assert all(findings.cell(2, col).data_type != "f" for col in (1, 3, 5))
    if with_previous:
        changes = workbook["Changes"]
        assert [changes.cell(2, col).value for col in (2, 3, 4, 5)] == [3, 2, 30, 4]
        assert [changes.cell(3, col).value for col in (2, 3, 4, 5)] == [-3, -2, -30, -2]


def test_header_only_findings_and_unmatched_model(tmp_path):
    previous = tmp_path / "previous"
    previous.mkdir()
    (previous / "model-a.json").write_bytes((PREVIOUS / "model-a.json").read_bytes())
    output = tmp_path / "report.xlsx"
    build_report(str(CURRENT), str(output), str(previous))
    workbook = load_workbook(output)
    assert workbook["Findings"].max_row == 1
    assert [workbook["Changes"].cell(3, col).value for col in (2, 3, 4, 5)] == [None] * 4


def test_existing_output_is_preserved(tmp_path):
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"original")
    with pytest.raises(ValueError, match="already exists"):
        build_report(str(CURRENT), str(output))
    assert output.read_bytes() == b"original"


def test_bad_inputs_fail_clearly(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(ValueError, match="no JSON files"):
        build_report(str(empty), str(tmp_path / "out.xlsx"))
    bad = tmp_path / "bad"
    bad.mkdir()
    data = json.loads((CURRENT / "model-a.json").read_text())
    data["schemaVersion"] = True
    (bad / "bad.json").write_text(json.dumps(data))
    with pytest.raises(ValueError, match=r"bad.json.*integer 1"):
        build_report(str(bad), str(tmp_path / "out.xlsx"))
    (bad / "bad.json").write_text("{")
    with pytest.raises(ValueError, match="bad.json"):
        build_report(str(bad), str(tmp_path / "out.xlsx"))
    (bad / "bad.json").write_text("[]")
    with pytest.raises(ValueError, match="root must be an object.*bad.json"):
        build_report(str(bad), str(tmp_path / "out.xlsx"))


def test_duplicate_titles_and_bad_findings(tmp_path):
    duplicate = tmp_path / "duplicate"
    duplicate.mkdir()
    for name in ("first", "second"):
        (duplicate / f"{name}.json").write_bytes((CURRENT / "model-a.json").read_bytes())
    with pytest.raises(ValueError, match="Duplicate current model title"):
        build_report(str(duplicate), str(tmp_path / "out.xlsx"), str(PREVIOUS))
    with pytest.raises(ValueError, match="Invalid finding 1"):
        build_report(
            str(CURRENT),
            str(tmp_path / "out.xlsx"),
            findings=[{"model": "A", "element_ids": [True]}],
        )


def test_mcp_is_local_and_accepts_camel_aliases(tmp_path):
    with patch.object(server, "_execute", side_effect=AssertionError("Revit contacted")):
        result = asyncio.run(
            server.mcp.call_tool(
                "revit_build_report",
                {
                    "snapshotsDir": str(CURRENT),
                    "outputPath": str(tmp_path / "report.xlsx"),
                    "findings": [
                        {
                            "model": "Model A",
                            "severity": "Low",
                            "rule": "Rule",
                            "elementIds": [7],
                            "recommendation": "Review",
                        }
                    ],
                },
            )
        )
    assert not result.is_error
    assert (tmp_path / "report.xlsx").exists()
    assert load_workbook(tmp_path / "report.xlsx")["Findings"]["D2"].value == "7"


def test_mcp_reports_expected_errors(tmp_path):
    with pytest.raises(ToolError, match="no JSON files"):
        empty = tmp_path / "empty"
        empty.mkdir()
        server.revit_build_report(str(empty), str(tmp_path / "out.xlsx"))
