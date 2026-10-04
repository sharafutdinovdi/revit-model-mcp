"""Validate review findings and write a local issue workbook."""

from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITIES = ("critical", "major", "minor", "info")
COLORS = ("C00000", "FF8C00", "FFD966", "D9E1F2")
SNAPSHOT_LIMIT = 25
HEADERS = [
    "ID",
    "Category",
    "Requirement source",
    "Requirement",
    "Finding",
    "Recommendation",
    "Severity",
    "Status",
    "Responsible",
    "Due",
    "Elements (count)",
    "Element IDs",
    "Snapshot",
]


def validate_capture(element_ids, pixel_size=1600, padding_mm=1500, mode="3d"):
    if not isinstance(element_ids, list) or not 1 <= len(element_ids) <= 500:
        raise ValueError("element_ids must contain 1 to 500 IDs.")
    if any(type(value) is not int or value <= 0 for value in element_ids):
        raise ValueError("element_ids must contain positive integers.")
    if type(pixel_size) is not int or not 1 <= pixel_size <= 4000:
        raise ValueError("pixel_size must be an integer from 1 to 4000.")
    if type(padding_mm) not in (int, float) or not 0 <= padding_mm <= 20000:
        raise ValueError("padding_mm must be from 0 to 20000.")
    if mode not in ("3d", "plan"):
        raise ValueError("mode must be 3d or plan.")


def validate_register(output_path, project, issues, pixel_size=900):
    """Return a destination and normalized input before any Revit request."""
    if not isinstance(output_path, str) or Path(output_path).suffix.lower() != ".xlsx":
        raise ValueError("output_path must have a .xlsx extension.")
    target = Path(output_path).expanduser().absolute()
    if target.exists() or target.is_symlink():
        raise ValueError(f"output_path already exists: {target}")
    if not isinstance(project, dict):
        raise ValueError("project must be an object.")
    for field in ("name", "model", "reviewer", "client", "stage", "date"):
        if field in project and not isinstance(project[field], str):
            raise ValueError(f"project.{field} must be a string.")
    documents = project.get("documents", [])
    if not isinstance(documents, list):
        raise ValueError("project.documents must be a list.")
    for index, document in enumerate(documents):
        if not isinstance(document, dict):
            raise ValueError(f"project.documents[{index}] must be an object.")
        for field in ("title", "reference", "revision"):
            if field in document and not isinstance(document[field], str):
                raise ValueError(f"project.documents[{index}].{field} must be a string.")
    if type(pixel_size) is not int or not 1 <= pixel_size <= 4000:
        raise ValueError("pixel_size must be an integer from 1 to 4000.")
    if not isinstance(issues, list) or not 1 <= len(issues) <= 60:
        raise ValueError("issues must contain 1 to 60 objects.")

    def check_text(value, field):
        if isinstance(value, str) and (ILLEGAL_CHARACTERS_RE.search(value) or len(value) > 32767):
            raise ValueError(
                f"{field} contains unsupported Excel text or exceeds 32767 characters."
            )

    for field, value in project.items():
        check_text(value, f"project.{field}")
    for index, document in enumerate(documents):
        for field, value in document.items():
            check_text(value, f"project.documents[{index}].{field}")
    normalized = []
    identifiers = set()
    for index, issue in enumerate(issues):
        prefix = f"issues[{index}]"
        if not isinstance(issue, dict):
            raise ValueError(f"{prefix} must be an object.")
        item = dict(issue)
        for field, value in item.items():
            check_text(value, f"{prefix}.{field}")
        for field in ("title", "category", "finding", "severity"):
            if not isinstance(item.get(field), str) or not item[field].strip():
                raise ValueError(f"{prefix}.{field} is required and must be a non-blank string.")
        for field in (
            "id",
            "requirement_source",
            "requirement",
            "recommendation",
            "status",
            "responsible",
            "due",
        ):
            if field in item and not isinstance(item[field], str):
                raise ValueError(f"{prefix}.{field} must be a string.")
        if item["severity"] not in SEVERITIES:
            raise ValueError(f"{prefix}.severity must be critical, major, minor or info.")
        item.setdefault("id", f"ISS-{index + 1:03d}")
        if not item["id"].strip() or item["id"] in identifiers:
            raise ValueError(f"{prefix}.id must be non-blank and unique.")
        identifiers.add(item["id"])
        item.setdefault("status", "Open")
        ids = item.get("element_ids", [])
        if (
            not isinstance(ids, list)
            or len(ids) > 500
            or any(type(value) is not int or value <= 0 for value in ids)
        ):
            raise ValueError(f"{prefix}.element_ids must contain at most 500 positive integers.")
        item["element_ids"] = list(dict.fromkeys(ids))
        item.setdefault("snapshot", "3d" if ids else "none")
        if item["snapshot"] not in ("3d", "plan", "none"):
            raise ValueError(f"{prefix}.snapshot must be 3d, plan or none.")
        check_text(f"{item['title']}\n{item['finding']}", f"{prefix}.finding")
        normalized.append(item)
    return target, dict(project), normalized


def snapshot_indices(issues):
    """Prioritize captures by severity, preserving input order within each severity."""
    candidates = [
        index
        for index, issue in enumerate(issues)
        if issue["element_ids"] and issue["snapshot"] != "none"
    ]
    ordered = sorted(candidates, key=lambda index: SEVERITIES.index(issues[index]["severity"]))
    return ordered[:SNAPSHOT_LIMIT], ordered[SNAPSHOT_LIMIT:]


def _append(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _header(sheet, widths):
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(bold=True, color="FFFFFF")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_register(
    output_path: str,
    project: dict[str, Any],
    issues: list[dict[str, Any]],
    snapshots: dict[int, str | Path] | None = None,
    snapshot_notes: dict[int, str] | None = None,
    warnings: list[str] | None = None,
) -> dict[str, Any]:
    """Write a new workbook. Snapshot paths remain readable until save completes."""
    target, project, issues = validate_register(output_path, project, issues)
    snapshots = snapshots or {}
    notes = dict(snapshot_notes or {})
    warnings = list(warnings or [])
    _, skipped = snapshot_indices(issues)
    for index in skipped:
        notes[index] = "Snapshot skipped: limit of 25 per register"
        warning = f"{issues[index]['id']}: {notes[index]}"
        if warning not in warnings:
            warnings.append(warning)
    for index, note in notes.items():
        warning = f"{issues[index]['id']}: {note}"
        if warning not in warnings:
            warnings.append(warning)
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    summary = workbook.create_sheet("Summary")
    register = workbook.create_sheet("Register")
    elements = workbook.create_sheet("Elements")
    workbook.properties.title = f"{project.get('name', '')} issue register".strip()
    workbook.properties.creator = project.get("reviewer", "")
    totals = {
        severity: sum(issue["severity"] == severity for issue in issues) for severity in SEVERITIES
    }
    _append(cover, ["Project field", "Value"])
    for field in ("name", "model", "reviewer", "client", "stage", "date"):
        _append(
            cover,
            [
                field.capitalize(),
                project.get(field, date.today().isoformat() if field == "date" else ""),
            ],
        )
    _append(cover, [])
    _append(cover, ["Document title", "Reference", "Revision"])
    for document in project.get("documents", []):
        _append(cover, [document.get(field, "") for field in ("title", "reference", "revision")])
    _append(cover, [])
    _append(cover, ["Severity", "Meaning", "Issues"])
    meanings = (
        "Blocks information exchange or coordination",
        "Breaks a contractual requirement",
        "Quality or consistency",
        "Observation",
    )
    for severity, color, meaning in zip(SEVERITIES, COLORS, meanings, strict=True):
        _append(cover, [severity, meaning, totals[severity]])
        cover.cell(cover.max_row, 1).fill = PatternFill("solid", fgColor=color)
        if severity == "critical":
            cover.cell(cover.max_row, 1).font = Font(color="FFFFFF")
    _header(cover, [28, 65, 20])
    _append(summary, ["Category", *[value.capitalize() for value in SEVERITIES], "Total"])
    categories = list(dict.fromkeys(issue["category"] for issue in issues))
    for category in categories:
        counts = [
            sum(issue["category"] == category and issue["severity"] == severity for issue in issues)
            for severity in SEVERITIES
        ]
        _append(summary, [category, *counts, sum(counts)])
    _append(summary, ["Total", *totals.values(), len(issues)])
    _append(summary, [])
    _append(summary, ["Status", "Issues"])
    for status, count in Counter(issue["status"] for issue in issues).items():
        _append(summary, [status, count])
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Issues by category and severity"
    chart.add_data(
        Reference(summary, min_col=2, max_col=5, min_row=1, max_row=len(categories) + 1),
        titles_from_data=True,
    )
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=len(categories) + 1))
    for series, color in zip(chart.series, COLORS, strict=True):
        series.graphicalProperties.solidFill = color
    summary.add_chart(chart, "H2")
    _header(summary, [35, 15, 15, 15, 15, 15])
    _append(register, HEADERS)
    _append(elements, ["Issue ID", "Element ID"])
    snapshot_count = 0
    for index, issue in enumerate(issues):
        ids = issue["element_ids"]
        _append(
            register,
            [
                issue["id"],
                issue["category"],
                issue.get("requirement_source", ""),
                issue.get("requirement", ""),
                f"{issue['title']}\n{issue['finding']}",
                issue.get("recommendation", ""),
                issue["severity"],
                issue["status"],
                issue.get("responsible", ""),
                issue.get("due", ""),
                len(ids),
                ", ".join(map(str, ids)),
                notes.get(index, ""),
            ],
        )
        row = register.max_row
        severity_cell = register.cell(row, 7)
        severity_cell.fill = PatternFill(
            "solid", fgColor=COLORS[SEVERITIES.index(issue["severity"])]
        )
        if issue["severity"] == "critical":
            severity_cell.font = Font(color="FFFFFF")
        register.row_dimensions[row].height = 70
        if index in snapshots and index not in skipped:
            try:
                image = Image(str(snapshots[index]))
            except (OSError, ValueError) as error:
                message = f"Snapshot unavailable: {error}"
                register.cell(row, 13, message)
                warnings.append(f"{issue['id']}: {message}")
                image = None
            if image is not None:
                image.height = image.height * 320 / image.width
                image.width = 320
                register.add_image(image, f"M{row}")
                register.row_dimensions[row].height = max(70, image.height * 0.75 + 8)
                snapshot_count += 1
        elif ids and issue["snapshot"] != "none" and index not in notes:
            message = "Snapshot unavailable: no image supplied"
            register.cell(row, 13, message)
            warnings.append(f"{issue['id']}: {message}")
        for element_id in ids:
            _append(elements, [issue["id"], element_id])
    _header(register, [16, 24, 24, 45, 55, 45, 14, 16, 24, 16, 18, 35, 48])
    _header(elements, [20, 20])
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        with target.open("xb") as output:
            workbook.save(output)
    except FileExistsError as error:
        raise ValueError(f"output_path already exists: {target}") from error
    return {
        "path": str(target),
        "issueCount": len(issues),
        "snapshotCount": snapshot_count,
        "bySeverity": totals,
        "warnings": warnings,
    }
