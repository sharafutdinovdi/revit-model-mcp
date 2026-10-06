"""Write model health counts and warning evidence to a local workbook."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from revit_model_mcp.atomic_write import write_new_file
from revit_model_mcp.workbook_style import (
    ACCENT,
    _append,
    _box_image,
    _cover_block,
    _header,
    _snapshot,
    _tile,
)

CHECKS = (
    ("Warnings", "warnings", 0, "Resolve model conflicts before delivery."),
    (
        "Views not on sheets",
        "viewsNotOnSheets",
        0.2,
        "Review undocumented views and model clutter.",
    ),
    ("Unused family types", "familyTypesUnused", 300, "Review unused content and model size."),
    ("In-place families", "familiesInPlace", 0, "Review content reuse and maintenance."),
    ("CAD imports", "cadImports", 0, "Review imported geometry and model performance."),
    ("Unplaced rooms", "roomsUnplaced", 0, "Review incomplete room placement."),
    ("Rooms not enclosed", "roomsNotEnclosed", 0, "Review room boundaries and area accuracy."),
    ("Design options", "designOptions", 0, "Review remaining alternatives at this review stage."),
    ("Images", "images", 20, "Review image content and model size."),
    ("Model groups", "groupsModel", 50, "Review group complexity and maintenance."),
)


def validate_health_path(save_to):
    """Validate the destination before contacting Revit and again before writing."""
    if not isinstance(save_to, str) or Path(save_to).suffix.lower() != ".xlsx":
        raise ValueError("save_to must have a .xlsx extension.")
    target = Path(save_to).expanduser().absolute()
    if target.exists() or target.is_symlink():
        raise ValueError(f"save_to already exists: {target}")
    if not target.parent.is_dir():
        raise ValueError("save_to parent directory must exist.")
    return target


def warning_ids(group):
    return list(dict.fromkeys(element["id"] for element in group.get("elements", [])))


def write_health_workbook(save_to, health, groups, snapshots=None, notes=None, warnings=None):
    """Keep capture failures as warnings and refuse replacement of an existing file."""
    target = validate_health_path(save_to)
    snapshots = snapshots or {}
    notes = notes or {}
    warnings = list(warnings or [])
    data = health["data"]
    counts = data["counts"]
    project = data.get("projectInfo") or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Health"
    warning_sheet = workbook.create_sheet("Warnings")
    count_sheet = workbook.create_sheet("Counts")
    workbook.properties.title = "Model health check"
    workbook.properties.creator = project.get("author") or ""
    _cover_block(sheet, 1, 1, 6, "Model health check", ACCENT, 22, True, end_row=2)
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 28
    for column in range(1, 7):
        sheet.cell(2, column).fill = PatternFill("solid", fgColor=ACCENT)
    size = data.get("fileSizeBytes")
    file_size = f"{size / 1024**2:.1f} MB" if size is not None else "File size unavailable"
    metadata = (
        f"{data.get('fileName') or 'Unknown file'} | Revit {data.get('revitVersion') or 'Unknown'}"
        f" | {file_size}"
    )
    _cover_block(sheet, 3, 1, 6, f"{metadata} | {date.today().isoformat()}", ACCENT, 12)
    sheet.row_dimensions[3].height = 30
    _cover_block(sheet, 5, 1, 6, "Project", size=12, bold=True)
    for row, field in enumerate(("name", "number", "client", "status", "author"), 6):
        _cover_block(sheet, row, 1, 1, field.capitalize(), "F2F2F2", bold=True)
        _cover_block(sheet, row, 2, 6, project.get(field) or "")
        sheet.row_dimensions[row].height = 24
    for column, (label, key) in enumerate(
        (
            ("Elements", "elements"),
            ("Warnings", "warnings"),
            ("Views", "views"),
            ("Sheets", "sheets"),
            ("Families", "families"),
            ("Links", "linksRvt"),
        ),
        1,
    ):
        value = counts.get(key)
        if label == "Links":
            cad = counts.get("linksCad")
            value = value + cad if value is not None and cad is not None else None
        _tile(
            sheet, 12, column, column, value if value is not None else "Unavailable", label, ACCENT
        )
    for column, value in enumerate(("Check", "Value", "Threshold", "Status", "Why it matters"), 1):
        sheet.cell(16, column, value)
    _header(sheet, [28, 20, 20, 16, 58, 18], row=16, columns=5)
    passed = 0
    for row, (label, key, threshold, reason) in enumerate(CHECKS, 17):
        value = counts.get(key)
        if key == "viewsNotOnSheets":
            views = counts.get("views")
            value = (
                (value / views if views else 0) if value is not None and views is not None else None
            )
        ok = value is not None and value <= threshold
        passed += ok
        _append(
            sheet,
            [
                label,
                value if value is not None else "Unavailable",
                threshold,
                "Pass" if ok else "Review",
                reason,
            ],
        )
        if key == "viewsNotOnSheets":
            sheet.cell(row, 2).number_format = "0.0%"
            sheet.cell(row, 3).number_format = "0%"
        status = sheet.cell(row, 4)
        status.fill = PatternFill("solid", fgColor="C6EFCE" if ok else "FFEB9C")
        status.font = Font(name="Calibri", size=10, bold=True, color="006100" if ok else "9C6500")
        sheet.row_dimensions[row].height = 32
    _append(sheet, ["Overall", f"{passed} of {len(CHECKS)} checks pass"])
    sheet.freeze_panes = "A17"
    sheet.print_title_rows = "16:16"
    _append(warning_sheet, ["Count", "Warning text", "Element IDs", "Snapshot"])
    snapshot_count = 0
    top = sorted(range(len(groups)), key=lambda index: groups[index]["count"], reverse=True)[:5]
    for index, group in enumerate(groups):
        ids = warning_ids(group)
        _append(
            warning_sheet,
            [
                group["count"],
                group["text"],
                ", ".join(map(str, ids[:30]))
                + (f" +{len(ids) - 30} more" if len(ids) > 30 else ""),
                notes.get(index, ""),
            ],
        )
        row = warning_sheet.max_row
        warning_sheet.row_dimensions[row].height = 60
        if index in snapshots and index in top:
            try:
                _box_image(warning_sheet, _snapshot(snapshots[index]), row, 4)
                snapshot_count += 1
            except (OSError, ValueError) as error:
                message = f"Snapshot unavailable: {error}"
                warning_sheet.cell(row, 4, message)
                warnings.append(f"Warning group {index + 1}: {message}")
    _header(warning_sheet, [12, 65, 32, 52], filter_rows=True)
    if not groups:
        _append(warning_sheet, [0, "No model warnings. Continue with the Health checks."])
    else:
        # Separate compact chart data from tall evidence rows.
        warning_sheet.cell(1, 6, "Warning text")
        warning_sheet.cell(1, 7, "Count")
        for row, index in enumerate(
            sorted(range(len(groups)), key=lambda i: groups[i]["count"], reverse=True)[:10], 2
        ):
            cell = warning_sheet.cell(row, 6, groups[index]["text"])
            cell.data_type = "s"
            warning_sheet.cell(row, 7, groups[index]["count"])
        warning_sheet.column_dimensions["F"].hidden = True
        warning_sheet.column_dimensions["G"].hidden = True
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Top warning groups"
        chart.legend = None
        chart.width = 22
        chart.height = 14
        chart.visible_cells_only = False
        chart.add_data(
            Reference(warning_sheet, min_col=7, min_row=1, max_row=min(10, len(groups)) + 1),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(warning_sheet, min_col=6, min_row=2, max_row=min(10, len(groups)) + 1)
        )
        warning_sheet.add_chart(chart, "I2")
    _append(count_sheet, ["Metric", "Value"])
    for metric, value in counts.items():
        _append(count_sheet, [metric, value if value is not None else "Unavailable"])
    _header(count_sheet, [32, 20], filter_rows=True)
    for tab in workbook:
        tab.sheet_view.showGridLines = False
        tab.page_setup.orientation = "landscape"
        tab.page_setup.paperSize = tab.PAPERSIZE_A3 if tab == warning_sheet else tab.PAPERSIZE_A4
        tab.page_setup.fitToWidth = 1
        tab.page_setup.fitToHeight = 0
        tab.sheet_properties.pageSetUpPr.fitToPage = True
        tab.print_options.horizontalCentered = True
        tab.oddFooter.left.text = (project.get("name") or "").replace("&", "&&")
        tab.oddFooter.right.text = "Page &P of &N"
        tab.print_area = tab.dimensions
        if tab != sheet:
            tab.print_title_rows = "1:1"
    warning_sheet.print_area = f"A1:U{max(30, warning_sheet.max_row)}"
    try:
        write_new_file(target, workbook.save)
    except ValueError as error:
        if target.exists() or target.is_symlink():
            raise ValueError(f"save_to already exists: {target}") from error
        raise
    except OSError as error:
        raise ValueError(f"Cannot save report: {error}") from error
    return {"path": str(target), "snapshotCount": snapshot_count, "warnings": warnings}
