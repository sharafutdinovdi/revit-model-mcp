"""Build a local Excel report from schema-version-1 model snapshots."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError

from revit_model_mcp.atomic_write import write_new_file

HEADERS = {
    "Summary": [
        "Model",
        "Saved-in year",
        "Runtime year",
        "Upgraded",
        "Workshared",
        "Number of saves",
        "File system time",
        "Revit Server last change",
        "Revit Server last user",
        "Elements",
        "Views",
        "Sheets",
        "Families",
        "Family types",
        "Links",
        "Warnings total",
        "Families total",
        "In-place families",
        "Family signals",
        "Skipped",
    ],
    "Warnings": ["Model", "Warning", "Count", "Element IDs"],
    "Families": [
        "Model",
        "Family",
        "Category",
        "In-place",
        "Editable",
        "Type count",
        "Instance count",
        "Warning count",
        "Signals",
    ],
    "Parameters": [
        "Model",
        "Category",
        "Parameter",
        "Total",
        "Filled",
        "Empty",
        "Missing",
        "Fill percent",
        "Sample IDs",
    ],
    "Skipped": ["Model", "What", "Reason"],
    "Changes": [
        "Model",
        "Warnings delta",
        "Families delta",
        "Fill percent delta",
        "Number of saves delta",
    ],
    "Findings": ["Model", "Severity", "Rule", "Element IDs", "Recommendation"],
}


def load_snapshots(directory: str) -> list[tuple[Path, dict[str, Any]]]:
    path = Path(directory).expanduser()
    if not path.is_dir():
        raise ValueError(f"Snapshot directory does not exist or is not a directory: {path}")
    files = sorted(file for file in path.glob("*.json") if file.is_file() and not file.is_symlink())
    if not files:
        raise ValueError(f"Snapshot directory contains no JSON files: {path}")
    snapshots = []
    for file in files:
        try:
            snapshot = json.loads(file.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as error:
            raise ValueError(f"Invalid snapshot JSON in {file}: {error}") from error
        if not isinstance(snapshot, dict):
            raise ValueError(f"Snapshot root must be an object: {file}")
        version = snapshot.get("schemaVersion")
        if type(version) is not int or version != 1:
            raise ValueError(f"Unsupported schemaVersion in {file}: expected integer 1")
        snapshots.append((file, snapshot))
    return snapshots


def _ids(values: list[int], truncated: bool = False) -> str:
    result = ", ".join(str(value) for value in values)
    return f"{result} [truncated]" if truncated else result


def _fill_percent(snapshot: dict[str, Any]) -> float:
    rows = snapshot["parameterFill"]["rows"]
    total = sum(row["total"] for row in rows)
    return sum(row["filled"] for row in rows) / total * 100 if total else 0


def _unique_titles(
    snapshots: list[tuple[Path, dict[str, Any]]], label: str
) -> dict[str, dict[str, Any]]:
    result = {}
    for file, snapshot in snapshots:
        try:
            title = snapshot["passport"]["title"]
            if not isinstance(title, str):
                raise TypeError("passport.title must be a string")
        except (KeyError, TypeError) as error:
            raise ValueError(f"Invalid {label} snapshot {file}: {error}") from error
        if title in result:
            raise ValueError(f"Duplicate {label} model title {title!r} in {file}")
        result[title] = snapshot
    return result


def _append(sheet: Any, values: list[Any]) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _finish(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25476A")
    for column in sheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(width + 2, 12), 55
        )


def build_report(
    snapshots_dir: str,
    output_path: str,
    previous_dir: str | None = None,
    findings: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Read local snapshots and create a new workbook without replacing a file."""
    current = load_snapshots(snapshots_dir)
    previous = load_snapshots(previous_dir) if previous_dir is not None else None
    prior_by_title = _unique_titles(previous, "previous") if previous is not None else {}
    if previous is not None:
        _unique_titles(current, "current")
    if not isinstance(output_path, str) or Path(output_path).suffix.lower() != ".xlsx":
        raise ValueError("Output path must have a .xlsx extension.")
    if findings is not None and not isinstance(findings, list):
        raise ValueError("Findings must be a list of objects.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    names = ["Summary", "Warnings", "Families", "Parameters", "Skipped"]
    if previous is not None:
        names.append("Changes")
    names.append("Findings")
    sheets = {name: workbook.create_sheet(name) for name in names}
    for name, sheet in sheets.items():
        _append(sheet, HEADERS[name])

    for file, snapshot in current:
        try:
            passport = snapshot["passport"]
            title = passport["title"]
            source = snapshot["source"]
            counts = passport["counts"]
            server = passport.get("revitServer") or {}
            families = snapshot["families"]
            _append(
                sheets["Summary"],
                [
                    title,
                    source["savedInYear"],
                    source["runtimeYear"],
                    source["upgradedInMemory"],
                    passport["isWorkshared"],
                    passport["numberOfSaves"],
                    passport["fileLastWriteUtc"],
                    server.get("lastModifiedUtc"),
                    server.get("lastModifiedBy"),
                    counts["elements"],
                    counts["views"],
                    counts["sheets"],
                    counts["families"],
                    counts["familyTypes"],
                    counts["links"],
                    snapshot["warnings"]["total"],
                    families["total"],
                    families["inPlace"],
                    len(families["signals"]),
                    snapshot["skippedCount"],
                ],
            )
            for group in snapshot["warnings"]["groups"]:
                _append(
                    sheets["Warnings"],
                    [
                        title,
                        group["text"],
                        group["count"],
                        _ids(group["elementIds"], group["elementIdsTruncated"]),
                    ],
                )
            for item in families["items"]:
                signals = sorted(
                    signal["signal"]
                    for signal in families["signals"]
                    if signal["family"] == item["name"]
                )
                _append(
                    sheets["Families"],
                    [
                        title,
                        item["name"],
                        item["category"],
                        item["isInPlace"],
                        item["isEditable"],
                        item["typeCount"],
                        item["instanceCount"],
                        item["warningCount"],
                        ", ".join(signals),
                    ],
                )
            for row in snapshot["parameterFill"]["rows"]:
                total = row["total"]
                _append(
                    sheets["Parameters"],
                    [
                        title,
                        row["category"],
                        row["parameter"],
                        total,
                        row["filled"],
                        row["empty"],
                        row["missing"],
                        row["filled"] / total * 100 if total else 0,
                        _ids(row["sampleEmptyIds"]),
                    ],
                )
            for skipped in snapshot["skipped"]:
                _append(sheets["Skipped"], [title, skipped["what"], skipped["reason"]])
            if previous is not None:
                prior = prior_by_title.get(title)
                _append(
                    sheets["Changes"],
                    [
                        title,
                        snapshot["warnings"]["total"] - prior["warnings"]["total"]
                        if prior
                        else None,
                        families["total"] - prior["families"]["total"] if prior else None,
                        _fill_percent(snapshot) - _fill_percent(prior) if prior else None,
                        passport["numberOfSaves"] - prior["passport"]["numberOfSaves"]
                        if prior
                        else None,
                    ],
                )
        except (KeyError, TypeError, ValueError, ZeroDivisionError, IllegalCharacterError) as error:
            raise ValueError(f"Invalid snapshot {file}: {error}") from error

    for index, finding in enumerate(findings or [], start=1):
        if not isinstance(finding, dict):
            raise ValueError(f"Finding {index} must be an object.")
        try:
            ids = finding.get("element_ids", finding.get("elementIds"))
            if not isinstance(ids, list) or any(type(value) is not int for value in ids):
                raise ValueError("element_ids must be a list of integers")
            values = [finding[key] for key in ("model", "severity", "rule")]
            values.extend([_ids(ids), finding["recommendation"]])
            if any(not isinstance(value, str) for value in (*values[:3], values[4])):
                raise ValueError("text fields must be strings")
            _append(sheets["Findings"], values)
        except (KeyError, ValueError, IllegalCharacterError) as error:
            raise ValueError(f"Invalid finding {index}: {error}") from error

    for sheet in sheets.values():
        _finish(sheet)
    target = Path(output_path).expanduser().absolute()
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target = target.parent.resolve() / target.name
        write_new_file(target, workbook.save)
    except ValueError as error:
        if target.exists() or target.is_symlink():
            raise ValueError(f"Local file already exists: {target}") from error
        raise
    except OSError as error:
        raise ValueError(f"Cannot save report: {error}") from error
    return {"outputPath": str(target), "snapshotCount": len(current), "sheets": names}
