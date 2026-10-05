"""Shared presentation helpers for local review workbooks."""

from io import BytesIO

from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PillowImage
from PIL import ImageChops, ImageOps

ACCENT = "1F3864"
COLORS = ("C00000", "FF8C00", "FFD966", "D9E1F2")
BORDER = Border(*(Side(style="thin", color="BFBFBF") for _ in range(4)))


def _append(sheet, values):
    sheet.append(values)
    if not values:
        return
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER
        if sheet.max_row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor="F2F2F2")


def _header(sheet, widths=None, row=1, columns=None, filter_rows=False):
    for cell in sheet[row][: columns or sheet.max_column]:
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    sheet.row_dimensions[row].height = 30
    for index, width in enumerate(widths or [], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if filter_rows:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions


def _cover_block(sheet, row, first, last, value, color=None, size=10, bold=False, end_row=None):
    sheet.merge_cells(start_row=row, start_column=first, end_row=end_row or row, end_column=last)
    cell = sheet.cell(row, first, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(
        name="Calibri",
        size=size,
        bold=bold,
        color="FFFFFF" if color in (ACCENT, COLORS[0]) else "404040",
    )
    cell.alignment = Alignment(
        wrap_text=True, vertical="center", horizontal="center" if color else "left"
    )
    if color:
        for column in range(first, last + 1):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=color)


def _snapshot(path):
    with PillowImage.open(path) as source:
        image = source.convert("RGBA")
        white = PillowImage.new("RGBA", image.size, "white")
        image = PillowImage.alpha_composite(white, image).convert("RGB")
    difference = ImageChops.difference(image, PillowImage.new("RGB", image.size, "white"))
    bounds = difference.convert("L").point(lambda value: 255 if value > 10 else 0).getbbox()
    if bounds:
        left, top, right, bottom = bounds
        image = image.crop(
            (
                max(0, left - 12),
                max(0, top - 12),
                min(image.width, right + 12),
                min(image.height, bottom + 12),
            )
        )
    image = ImageOps.contain(image, (360, 240), PillowImage.Resampling.LANCZOS)
    canvas = PillowImage.new("RGB", (360, 240), "white")
    canvas.paste(image, ((360 - image.width) // 2, (240 - image.height) // 2))
    buffer = BytesIO()
    canvas.save(buffer, format="PNG")
    buffer.seek(0)
    return Image(buffer)


def _tile(sheet, row, first, last, value, label, color):
    _cover_block(sheet, row, first, last, value, color, 20, True)
    _cover_block(sheet, row + 1, first, last, label, color)
    sheet.row_dimensions[row].height = 36
    sheet.row_dimensions[row + 1].height = 24


def _box_image(sheet, image, row, column):
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=column - 1, row=row - 1, colOff=pixels_to_EMU(4), rowOff=pixels_to_EMU(3)
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(360), pixels_to_EMU(240)),
    )
    sheet.add_image(image)
    sheet.row_dimensions[row].height = 185
